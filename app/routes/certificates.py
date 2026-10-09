
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    Form,
    HTTPException,
    UploadFile,
)
from fastapi.responses import FileResponse
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.database import SessionLocal, get_db
from app.models import Certificate, GenerationJob, Status
from app.schemas import (
    CertificateResponse,
    CertificateResultList,
    GenerationJobInput,
    GenerationJobResponse,
    JobStatusResponse,
    RecipientInput,
)
from app.services.certificate_generator import generate_certificate


router = APIRouter(prefix="/api", tags=["Certificates"])

MAX_RECIPIENTS = 500
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
COMMIT_EVERY = 10


# ---------- Helpers ----------

def get_job_or_404(db: Session, job_id: int) -> GenerationJob:
    job = db.get(GenerationJob, job_id)

    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")

    return job


def create_job_record(
    db: Session,
    data: GenerationJobInput,
    background_tasks: BackgroundTasks,
) -> GenerationJobResponse:
    """Create one job and its certificate records."""

    total = len(data.recipients)

    job = GenerationJob(
        event_name=data.event_name,
        event_date=data.event_date,
        organization=data.organization,
        total_count=total,
    )

    try:
        db.add(job)
        db.flush()

        job_id = job.id

        db.add_all(
            [
                Certificate(
                    job_id=job_id,
                    recipient_name=recipient.name,
                    recipient_email=str(recipient.email),
                )
                for recipient in data.recipients
            ]
        )

        db.commit()

    except Exception:
        db.rollback()
        raise

    background_tasks.add_task(process_generation_job, job_id)

    return GenerationJobResponse(
        job_id=job_id,
        status=Status.PENDING,
        total_recipients=total,
    )


def read_recipients_from_excel(
    file_contents: bytes,
) -> list[RecipientInput]:
    """Read and validate recipients from an Excel workbook."""

    try:
        workbook = load_workbook(
            filename=BytesIO(file_contents),
            read_only=True,
            data_only=True,
        )
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid or corrupted Excel file.",
        )

    try:
        rows = workbook.active.iter_rows(values_only=True)
        header_row = next(rows, None)

        if header_row is None:
            raise HTTPException(
                status_code=400,
                detail="Excel file is empty.",
            )

        headers = [
            str(cell).strip().lower() if cell is not None else ""
            for cell in header_row
        ]

        if "name" not in headers or "email" not in headers:
            raise HTTPException(
                status_code=400,
                detail="Excel file must contain Name and Email columns.",
            )

        name_index = headers.index("name")
        email_index = headers.index("email")

        recipients: list[RecipientInput] = []

        for row_number, row in enumerate(rows, start=2):
            name = row[name_index] if name_index < len(row) else None
            email = row[email_index] if email_index < len(row) else None

            # Skip completely empty rows.
            if name is None and email is None:
                continue

            try:
                recipient = RecipientInput(
                    name=str(name).strip() if name is not None else "",
                    email=str(email).strip() if email is not None else "",
                )
            except ValidationError:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"Invalid recipient data in Excel row {row_number}."
                    ),
                )

            recipients.append(recipient)

            if len(recipients) > MAX_RECIPIENTS:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"A maximum of {MAX_RECIPIENTS} "
                        "recipients is allowed per job."
                    ),
                )

        if not recipients:
            raise HTTPException(
                status_code=400,
                detail="Excel file contains no valid recipient rows.",
            )

        return recipients

    finally:
        workbook.close()


# ---------- Background worker ----------

def process_generation_job(job_id: int) -> None:
    """Generate certificates and track individual results."""

    db = SessionLocal()

    try:
        job = db.get(GenerationJob, job_id)

        if job is None:
            return

        job.status = Status.PROCESSING
        db.commit()

        certificates = (
            db.query(Certificate)
            .filter(
                Certificate.job_id == job_id,
                Certificate.status == Status.PENDING,
            )
            .order_by(Certificate.id)
            .all()
        )

        event_name = job.event_name
        event_date = job.event_date.isoformat()
        organization = job.organization

        for index, certificate in enumerate(certificates, start=1):
            try:
                certificate.file_path = generate_certificate(
                    certificate_id=certificate.id,
                    recipient_name=certificate.recipient_name,
                    event_name=event_name,
                    event_date=event_date,
                    organization=organization,
                )

                certificate.status = Status.SUCCESS
                certificate.error_message = None
                job.success_count += 1

            except Exception:
                # Avoid exposing internal exception details to API users.
                certificate.status = Status.FAILED
                certificate.error_message = (
                    "Certificate generation failed."
                )
                job.failed_count += 1

            if index % COMMIT_EVERY == 0:
                db.commit()

        job.status = (
            Status.COMPLETED_WITH_ERRORS
            if job.failed_count
            else Status.COMPLETED
        )
        job.completed_at = datetime.now(timezone.utc)
        db.commit()

    except Exception:
        db.rollback()

        try:
            job = db.get(GenerationJob, job_id)

            if job is not None:
                job.status = Status.FAILED
                job.completed_at = datetime.now(timezone.utc)
                db.commit()

        except Exception:
            db.rollback()

    finally:
        db.close()


# ---------- Routes ----------

@router.post(
    "/jobs",
    response_model=GenerationJobResponse,
    status_code=202,
)
def create_generation_job(
    request: GenerationJobInput,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
):
    """Create a bulk certificate generation job using JSON."""

    return create_job_record(db, request, background_tasks)


@router.post(
    "/jobs/upload",
    response_model=GenerationJobResponse,
    status_code=202,
)
def create_job_from_excel(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    event_name: str = Form(...),
    event_date: date = Form(...),
    organization: str = Form(...),
    db: Session = Depends(get_db),
):
    """Create a bulk job from an Excel file."""

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Please upload an .xlsx Excel file.",
        )

    file_contents = file.file.read(MAX_UPLOAD_BYTES + 1)

    if len(file_contents) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File too large (maximum 5 MB).",
        )

    recipients = read_recipients_from_excel(file_contents)

    try:
        job_input = GenerationJobInput(
            event_name=event_name,
            event_date=event_date,
            organization=organization,
            recipients=recipients,
        )
    except ValidationError as exc:
        raise HTTPException(
            status_code=422,
            detail=exc.errors(),
        )

    return create_job_record(db, job_input, background_tasks)


@router.get(
    "/jobs/{job_id}",
    response_model=JobStatusResponse,
)
def get_job_status(
    job_id: int,
    db: Session = Depends(get_db),
):
    """Return the overall job status and progress."""

    job = get_job_or_404(db, job_id)

    return JobStatusResponse(
        job_id=job.id,
        status=job.status,
        total=job.total_count,
        successful=job.success_count,
        failed=job.failed_count,
        completed=job.success_count + job.failed_count,
    )


@router.get(
    "/jobs/{job_id}/certificates",
    response_model=CertificateResultList,
)
def list_job_certificates(
    job_id: int,
    db: Session = Depends(get_db),
):
    """List the individual certificate results for a job."""

    job = get_job_or_404(db, job_id)

    certificates = (
        db.query(Certificate)
        .filter(Certificate.job_id == job_id)
        .order_by(Certificate.id)
        .all()
    )

    return CertificateResultList(
        job_id=job.id,
        certificates=[
            CertificateResponse(
                certificate_id=certificate.id,
                recipient_name=certificate.recipient_name,
                recipient_email=certificate.recipient_email,
                status=certificate.status,
                error_message=certificate.error_message,
            )
            for certificate in certificates
        ],
    )


@router.get("/certificates/{certificate_id}")
def download_certificate(
    certificate_id: int,
    db: Session = Depends(get_db),
):

    certificate = db.get(Certificate, certificate_id)

    if certificate is None:
        raise HTTPException(
            status_code=404,
            detail="Certificate not found.",
        )

    if (
        certificate.status != Status.SUCCESS
        or not certificate.file_path
    ):
        raise HTTPException(
            status_code=409,
            detail="Certificate is not available for download.",
        )

    file_path = Path(certificate.file_path)

    if not file_path.is_file():
        raise HTTPException(
            status_code=404,
            detail="Certificate file not found.",
        )

    safe_recipient = certificate.recipient_name.strip().replace(" ", "_")
    safe_event = certificate.job.event_name.strip().replace(" ", "_")

    return FileResponse(
        path=file_path,
        media_type="application/pdf",
        filename=f"{safe_recipient}_{safe_event}.pdf"
    )